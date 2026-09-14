from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"C:\Users\roiko\PycharmProjects\oks-labs\lab1\Отчет_ЛР1_Ройко.docx"


def set_run_font(run, name="Times New Roman", size=14, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def setup_paragraph(
    paragraph,
    align="justify",
    first_line=True,
    space_after=6,
    space_before=0,
    line_spacing=1.5,
):
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.widow_control = True
    if align == "justify":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
    else:
        paragraph.paragraph_format.first_line_indent = Cm(0)


def body(doc, text):
    p = doc.add_paragraph()
    setup_paragraph(p)
    add_text(p, text)
    return p


def heading(doc, text, center=True):
    p = doc.add_paragraph()
    setup_paragraph(p, align="center" if center else "left", first_line=False, space_before=12, space_after=10)
    add_text(p, text, bold=True)
    return p


def subheading(doc, text):
    p = doc.add_paragraph()
    setup_paragraph(p, align="left", first_line=False, space_before=10, space_after=8)
    add_text(p, text, bold=True)
    return p


def center_line(doc, text, bold=False, size=14, space_after=0, space_before=0):
    p = doc.add_paragraph()
    setup_paragraph(p, align="center", first_line=False, space_after=space_after, space_before=space_before)
    add_text(p, text, bold=bold, size=size)
    return p


def empty(doc, n=1):
    for _ in range(n):
        p = doc.add_paragraph()
        setup_paragraph(p, align="center", first_line=False, space_after=0, space_before=0, line_spacing=1.0)
        add_text(p, "")


def add_code(doc, code):
    for line in code.split("\n"):
        p = doc.add_paragraph()
        setup_paragraph(p, align="left", first_line=False, space_after=0, space_before=0, line_spacing=1.15)
        p.paragraph_format.left_indent = Cm(1.0)
        add_text(p, line if line else " ", name="Courier New", size=11)


def figure_caption(doc, text):
    p = doc.add_paragraph()
    setup_paragraph(p, align="center", first_line=False, space_before=6, space_after=12)
    add_text(p, text, italic=True)


def set_narrow_margins(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)


def prevent_widow(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    keep.set(qn("w:val"), "true")


def build():
    doc = Document()
    set_narrow_margins(doc.sections[0])
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # Title page
    center_line(doc, "Министерство образования Республики Беларусь")
    center_line(doc, "Учреждение образования")
    center_line(doc, "БЕЛОРУССКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ", bold=True)
    center_line(doc, "ИНФОРМАТИКИ И РАДИОЭЛЕКТРОНИКИ", bold=True)
    empty(doc, 2)
    center_line(doc, "Факультет компьютерных систем и сетей")
    center_line(doc, "Кафедра электронных вычислительных машин")
    empty(doc, 4)
    center_line(doc, "Лабораторная работа №1", bold=True)
    center_line(doc, "по дисциплине «Основы компьютерных сетей»", bold=True)
    center_line(doc, "«COM-порты устройств»", bold=True)
    empty(doc, 6)

    p = doc.add_paragraph()
    setup_paragraph(p, align="left", first_line=False, space_after=4)
    p.paragraph_format.left_indent = Cm(9.0)
    add_text(p, "Студент гр. 250502")
    tab = p.add_run("\t")
    set_run_font(tab)
    add_text(p, "Ройко Артём")

    p = doc.add_paragraph()
    setup_paragraph(p, align="left", first_line=False, space_after=4)
    p.paragraph_format.left_indent = Cm(9.0)
    add_text(p, "Преподаватель")
    add_text(p, "\t\t")
    add_text(p, "Глецевич И.И.")

    empty(doc, 8)
    center_line(doc, "МИНСК 2026")

    doc.add_page_break()

    heading(doc, "ЦЕЛЬ РАБОТЫ")
    body(
        doc,
        "Целью лабораторной работы является получение практических навыков "
        "работы с COM-портами и обучение программированию COM-портов.",
    )

    heading(doc, "1 ТЕОРЕТИЧЕСКАЯ ЧАСТЬ")

    subheading(doc, "1.1 Исходные данные")
    body(
        doc,
        "Лабораторная работа выполнена в составе бригады по варианту 2. "
        "Моя часть работы — графический интерфейс коммуникационной программы. "
        "Целевая операционная система — Windows 10. Язык программирования — Python 3. "
        "Графический интерфейс реализован библиотекой PyQt6. Для доступа к последовательному "
        "порту используется библиотека pyserial. Среда разработки — Visual Studio Code. "
        "Интерфейс собран программно, без Qt Designer.",
    )
    body(
        doc,
        "Программа предназначена для обмена текстовыми сообщениями между двумя "
        "пользовательскими станциями через COM-порты по предсозданной топологии "
        "лаборатории 509-V. Номер порта определяется в Device Manager (обычно COM3). "
        "Согласно варианту 2 выбираемым параметром UART 16550 является длина байта; "
        "остальные параметры порта фиксированы. Передача сообщений выполняется "
        "построчно после нажатия Enter. Язык элементов интерфейса — английский.",
    )

    subheading(doc, "1.2 Параметры инициализации COM-порта")
    body(
        doc,
        "Для работы с последовательным портом применяется класс serial.Serial "
        "библиотеки pyserial. Список имён портов, существующих в системе, возвращает "
        "функция serial.tools.list_ports.comports(). Конструктор Serial открывает порт "
        "и задаёт параметры связи в соответствии с архитектурой UART 16550. Ниже "
        "перечислены все основные параметры инициализации и все возможные значения, "
        "включая незадействованные в данной работе.",
    )
    body(
        doc,
        "port (str) — системное имя порта, например COM1, COM3.",
    )
    body(
        doc,
        "baudrate (int) — скорость передачи в бодах. Стандартные значения: 110, 300, "
        "600, 1200, 2400, 4800, 9600, 19200, 38400, 57600, 115200 и другие. "
        "В работе параметр незадействован для выбора пользователем и фиксирован: 9600.",
    )
    body(
        doc,
        "bytesize — число бит данных в одном кадре. Возможные значения: 5 (FIVEBITS), "
        "6 (SIXBITS), 7 (SEVENBITS), 8 (EIGHTBITS). В варианте 2 этот параметр "
        "выбирается пользователем в окне управления.",
    )
    body(
        doc,
        "parity — контроль чётности. Возможные значения: PARITY_NONE («N») — контроль "
        "отсутствует; PARITY_EVEN («E») — чётный; PARITY_ODD («O») — нечётный; "
        "PARITY_MARK («M») — бит паритета всегда 1; PARITY_SPACE («S») — бит паритета "
        "всегда 0. В работе параметр незадействован для выбора и фиксирован: PARITY_NONE.",
    )
    body(
        doc,
        "stopbits — число стоп-битов. Возможные значения: 1 (STOPBITS_ONE), "
        "1.5 (STOPBITS_ONE_POINT_FIVE), 2 (STOPBITS_TWO). Старт-бит всегда один. "
        "В работе параметр незадействован для выбора и фиксирован: 1.",
    )
    body(
        doc,
        "timeout и write_timeout (float либо None) — таймауты чтения и записи в секундах. "
        "None — блокирующий режим, 0 — неблокирующий. В работе: timeout = 0.05, "
        "write_timeout = 1.",
    )
    body(
        doc,
        "xonxoff, rtscts, dsrdtr (bool) — программное (XON/XOFF) и аппаратное "
        "(RTS/CTS, DSR/DTR) управление потоком. Возможные значения: True и False. "
        "В работе незадействованы и остаются False по умолчанию.",
    )
    body(
        doc,
        "Основные методы класса Serial: open(), close(), write(data), read(size). "
        "Основные свойства: is_open, in_waiting. В интерфейсе для выбора порта "
        "используется виджет QComboBox (класс HoverComboBox), для выбора длины байта — "
        "второй QComboBox, для ввода и вывода сообщений — QPlainTextEdit, для состояния — QLabel.",
    )

    heading(doc, "2 ПРАКТИЧЕСКАЯ ЧАСТЬ")
    body(
        doc,
        "В рамках своей части работы реализован графический интерфейс отдельного "
        "приложения на PyQt6. Интерфейс содержит четыре панели: Control — управление, "
        "State — просмотр состояния, Input — ввод сообщений для передачи, Output — "
        "вывод принятых сообщений. Назначение элементов соответствует общепринятому "
        "смыслу. Окно управления содержит только два элемента: форму выбора COM-порта "
        "с отображением выбранного значения и форму выбора длины байта со значениями "
        "5, 6, 7 и 8. Отладочное окно не добавлялось.",
    )
    body(
        doc,
        "Поле Input постоянно доступно для ввода до закрытия программы и принимает "
        "все печатные символы, а также Enter. Согласно варианту 2 передача строки "
        "привязана к нажатию Enter. Панель Output предназначена только для принятого "
        "текста и недоступна для редактирования. Формы соразмерны наполнению: "
        "Control и State компактны, Input и Output занимают основную площадь окна. "
        "Списки порта и длины байта после успешного открытия порта блокируются, "
        "что обеспечивает однократный выбор.",
    )

    subheading(doc, "2.1 Создание виджетов интерфейса")
    body(
        doc,
        "В методе create_widgets создаются выпадающие списки COM-порта и размера байта, "
        "поля ввода и вывода, а также метка состояния. Список порта заполняется именами, "
        "полученными из системы; выбранное значение отображается в самом списке. "
        "HoverComboBox открывает перечень значений при наведении указателя; стрелка "
        "списка является частью этого же элемента, а не отдельной кнопкой окна управления.",
    )
    add_code(
        doc,
        "def create_widgets(self):\n"
        "    self.port_combo = HoverComboBox()\n"
        "    self.byte_size_combo = HoverComboBox()\n"
        "    for byte_size in (5, 6, 7, 8):\n"
        "        self.byte_size_combo.addItem(str(byte_size), byte_size)\n"
        "    self.byte_size_combo.setCurrentIndex(-1)\n"
        "    self.input_text = MessageInput()\n"
        "    self.output_text = QtWidgets.QPlainTextEdit()\n"
        "    self.output_text.setReadOnly(True)\n"
        "    self.status_label = QtWidgets.QLabel()",
    )

    subheading(doc, "2.2 Компоновка панелей Control, State, Input и Output")
    body(
        doc,
        "В методе create_layout панели управления и состояния размещаются в верхней "
        "части окна, панели ввода и вывода — в нижней. Для группировки использован "
        "QFrame, заголовки панелей заданы подписями Control, State, Input и Output.",
    )
    add_code(
        doc,
        "def create_layout(self):\n"
        "    control_panel = self.create_panel(\"Control\")\n"
        "    status_panel = self.create_panel(\"State\")\n"
        "    input_group = self.create_panel(\"Input\")\n"
        "    output_group = self.create_panel(\"Output\")\n"
        "    control_grid.addWidget(self.port_combo, 1, 0)\n"
        "    control_grid.addWidget(self.byte_size_combo, 1, 1)\n"
        "    top_layout.addWidget(control_panel, 7)\n"
        "    top_layout.addWidget(status_panel, 3)\n"
        "    messages_layout.addWidget(input_group, 1)\n"
        "    messages_layout.addWidget(output_group, 1)",
    )

    subheading(doc, "2.3 Настройка стилей")
    body(
        doc,
        "В методе apply_styles через setStyleSheet задана единая тёмная тема: цвет фона "
        "окна, рамки панелей, оформление списков QComboBox и полей QPlainTextEdit. "
        "Все подписи интерфейса выполнены на одном языке — английском.",
    )
    add_code(
        doc,
        "def apply_styles(self):\n"
        "    self.setStyleSheet(\"\"\"\n"
        "        QWidget { background: #101418; color: #ecfeff; }\n"
        "        QFrame#panel { background: #182026; border: 1px solid #2f3d46; }\n"
        "        QComboBox, QPlainTextEdit { background: #0d1216; }\n"
        "    \"\"\")",
    )
    body(
        doc,
        "В панель State периодически, один раз в секунду, выводится количество "
        "переданных символов в виде строки Sent characters: N. Ошибки открытия порта, "
        "отправки и приёма отображаются пользователю отдельным диалогом Error и "
        "не смешиваются с отладочными сообщениями.",
    )

    subheading(doc, "2.4 Результат")
    body(
        doc,
        "Результат работы интерфейса показан на рисунке 1. Рисунок выполнен упрощённо, "
        "но отражает состав панелей и двух элементов окна управления.",
    )
    add_code(
        doc,
        "+--------------------------------------------------+\n"
        "| Lab 1: COM Port                                  |\n"
        "| +----------------------+  +-------------------+  |\n"
        "| | Control              |  | State             |  |\n"
        "| | COM port  [ ......v] |  | Sent characters: 0|  |\n"
        "| | Byte size [ ......v] |  |                   |  |\n"
        "| +----------------------+  +-------------------+  |\n"
        "| +----------------------+  +-------------------+  |\n"
        "| | Input                |  | Output            |  |\n"
        "| |                      |  |                   |  |\n"
        "| +----------------------+  +-------------------+  |\n"
        "+--------------------------------------------------+",
    )
    figure_caption(doc, "Рисунок 1 — Главное окно программы (упрощённая схема)")

    heading(doc, "3 ЗАКЛЮЧЕНИЕ")
    body(
        doc,
        "В ходе лабораторной работы в составе бригады разработано приложение для "
        "обмена данными через COM-порт по варианту 2. Выбираемый параметр — длина "
        "байта (5–8); скорость 9600 бод, контроль паритета none и один стоп-бит "
        "фиксированы. Передача выполняется построчно после нажатия Enter, данные "
        "передаются посимвольно как сырой поток. Принятое сообщение отображается "
        "в окне вывода сразу по мере поступления.",
    )
    body(
        doc,
        "Моя часть работы — графический интерфейс на PyQt6. Созданы отдельные панели "
        "Control, State, Input и Output; в окне управления размещены только два "
        "элемента — список COM-порта и список длины байта; в окне состояния "
        "периодически выводится количество переданных символов; задана единая тёмная "
        "тема и английский язык элементов. Интерфейс постоянно готов к вводу символов "
        "до закрытия программы. Дополнительный функционал в состав интерфейса не включался.",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
